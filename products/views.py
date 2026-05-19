import csv
import json
import logging
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from datetime import timedelta
from io import BytesIO
from math import ceil
from django.apps import apps
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, HttpResponseBadRequest, HttpResponseForbidden
from django.shortcuts import render, redirect, get_object_or_404
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET

from branches.models import Branch
from .forms import (
    ProductForm,
    CategoryForm,
    AdjustStockForm,
    TransferStockForm,
    ProductBulkUploadForm,
    SupplierForm,
    SupplierReorderResponseForm,
)
from .models import (
    AutoReorderRequest,
    Category,
    Supplier,
    SupplierReorderRequest,
    Product,
    Stock,
    StockMovement,
    Purchase,
    PurchaseItem,
    Transfer,
    TransferItem,
)
from .tasks import notify_next_supplier, send_purchase_confirmation_to_supplier

logger = logging.getLogger(__name__)

BULK_UPLOAD_FAILED_ROWS_SESSION_KEY = "products_bulk_upload_failed_rows_report"


def _to_non_negative_int(value):
    try:
        return max(int(value), 0)
    except (TypeError, ValueError):
        return 0


def _branch_pending_packets(supplier_request, branch):
    if not branch:
        return supplier_request.pending_quantity

    branch_requirements = supplier_request.reorder_request.branch_requirements or {}
    if not isinstance(branch_requirements, dict) or not branch_requirements:
        return supplier_request.pending_quantity

    return _to_non_negative_int(branch_requirements.get(branch.name, 0))


def _normalized_branch_requirements(branch_requirements):
    if not isinstance(branch_requirements, dict):
        return []

    payload = []
    for branch_name, qty in branch_requirements.items():
        normalized_qty = _to_non_negative_int(qty)
        if normalized_qty > 0:
            payload.append((str(branch_name), normalized_qty))

    payload.sort(key=lambda item: item[0].lower())
    return payload


def _build_branch_groups(branch_product_map):
    groups = []
    for branch_name in sorted(branch_product_map.keys(), key=lambda value: value.lower()):
        branch_items = sorted(
            branch_product_map[branch_name],
            key=lambda item: item["product_name"].lower(),
        )
        groups.append(
            {
                "name": branch_name,
                "items": branch_items,
                "total_packets": sum(int(item.get("quantity", 0) or 0) for item in branch_items),
                "total_units": sum(
                    int(item.get("quantity_units", item.get("quantity", 0)) or 0) for item in branch_items
                ),
            }
        )
    return groups


def _allocate_confirmed_branch_quantities(branch_requirements, confirmed_quantity):
    remaining = _to_non_negative_int(confirmed_quantity)
    if remaining <= 0:
        return []

    allocations = []
    for branch_name, requested_qty in _normalized_branch_requirements(branch_requirements):
        if remaining <= 0:
            break
        allocated_qty = min(requested_qty, remaining)
        if allocated_qty <= 0:
            continue
        allocations.append((branch_name, allocated_qty))
        remaining -= allocated_qty

    if remaining > 0:
        allocations.append(("Unspecified Branch", remaining))

    return allocations


def _confirmed_branch_groups_for_supplier_request(supplier_request):
    confirmed_qty = _to_non_negative_int(supplier_request.fulfilled_quantity)
    if confirmed_qty <= 0:
        return []

    pack_quantity = max(int(supplier_request.reorder_request.product.pack_quantity or 1), 1)
    branch_product_map = {}
    for branch_name, branch_qty in _allocate_confirmed_branch_quantities(
        supplier_request.reorder_request.branch_requirements or {},
        confirmed_qty,
    ):
        branch_product_map.setdefault(branch_name, []).append(
            {
                "product_name": supplier_request.reorder_request.product.name,
                "quantity": branch_qty,
                "quantity_units": branch_qty * pack_quantity,
            }
        )
    return _build_branch_groups(branch_product_map)


def _used_supplier_link_message(supplier_request):
    if supplier_request.status in {
        SupplierReorderRequest.STATUS_ACCEPTED,
        SupplierReorderRequest.STATUS_PARTIAL,
        SupplierReorderRequest.STATUS_REJECTED,
    }:
        return "warning", "This confirmation link has already been used. It is now read-only for your records."
    if supplier_request.status == SupplierReorderRequest.STATUS_EXPIRED:
        return "warning", "This confirmation link has expired and can no longer accept responses."
    if supplier_request.status == SupplierReorderRequest.STATUS_EMAIL_FAILED:
        return "error", "This request failed to send correctly. Please contact our procurement team."
    return "info", "This confirmation link is no longer active."


def _max_packets_allowed_for_product(product):
    pack_quantity = max(int(product.pack_quantity or 1), 1)
    max_stock_units = max(int(product.max_stock or 1), 1)
    return max(int(ceil(max_stock_units / float(pack_quantity))), 1)


def _get_or_create_unregistered_supplier(unregistered_name: str) -> Supplier:
    canonical_name = "Unregistered Supplier"
    existing = Supplier.objects.filter(name__iexact=canonical_name).order_by("id").first()
    if existing:
        return existing

    return Supplier.objects.create(
        name=canonical_name,
        contact_person=canonical_name,
        phone_number="0000000000",
        email="unregistered-supplier@example.com",
        address="Unregistered supplier",
        priority=9999,
    )


def _pending_packets_by_product_branch(product_ids):
    totals = defaultdict(int)
    if not product_ids:
        return totals

    open_reorders = AutoReorderRequest.objects.filter(
        status=AutoReorderRequest.STATUS_OPEN,
        product_id__in=product_ids,
    ).values_list("product_id", "branch_requirements")

    for product_id, branch_requirements in open_reorders:
        if not isinstance(branch_requirements, dict):
            continue
        for branch_name, packets in branch_requirements.items():
            normalized_packets = _to_non_negative_int(packets)
            if normalized_packets <= 0:
                continue
            totals[(int(product_id), str(branch_name))] += normalized_packets

    return totals


@login_required
@require_GET
def supplier_pending_orders_view(request):
    supplier_id = request.GET.get("supplier_id")
    if not supplier_id:
        return HttpResponse("")

    supplier = get_object_or_404(Supplier, pk=supplier_id)
    active_branch = _resolve_products_branch(request)
    can_edit_quantity = _can_edit_receive_packets(request.user)
    pending_requests = []

    orders = SupplierReorderRequest.objects.filter(
        supplier=supplier,
        status__in=[SupplierReorderRequest.STATUS_ACCEPTED, SupplierReorderRequest.STATUS_PARTIAL]
    ).select_related("reorder_request", "reorder_request__product")

    for order in orders:
        if order.pending_quantity > 0:
            branch_pending_quantity = _branch_pending_packets(order, active_branch)
            if branch_pending_quantity <= 0:
                continue
            order.branch_pending_quantity = min(order.pending_quantity, branch_pending_quantity)
            pending_requests.append(order)

    return render(
        request,
        "products/partials/_pending_order_rows.html",
        {
            "pending_requests": pending_requests,
            "can_edit_quantity": can_edit_quantity,
        },
    )

@login_required
def manual_order_create_view(request):
    if not _can_manage_catalog(request.user):
        return HttpResponseForbidden("Permission denied.")

    Branch = apps.get_model("branches", "Branch")
    can_select_branch = request.user.can_manage_users()
    all_branches = Branch.objects.filter(is_active=True).order_by("name") if can_select_branch else None
    
    active_branch_id = request.GET.get("branch") or request.POST.get("branch_id")
    if not can_select_branch:
        active_branch = request.user.branch
        active_branch_id = active_branch.id if active_branch else None
    else:
        active_branch = Branch.objects.filter(pk=active_branch_id).first() if active_branch_id else None

    if request.method == "POST":
        supplier_strategy = (request.POST.get("supplier_strategy") or "cascade").strip()
        selected_supplier_ids_raw = request.POST.getlist("supplier_id[]")
        unregistered_supplier_name = (request.POST.get("unregistered_supplier_name") or "").strip()
        preferred_supplier_ids = []

        if supplier_strategy not in {"cascade", "selected", "unregistered"}:
            return HttpResponseBadRequest("Invalid supplier strategy selected.")

        if supplier_strategy == "selected":
            if not selected_supplier_ids_raw:
                return HttpResponseBadRequest("Select at least one supplier.")
            try:
                preferred_supplier_ids = sorted({int(value) for value in selected_supplier_ids_raw if str(value).strip()})
            except (TypeError, ValueError):
                return HttpResponseBadRequest("Invalid supplier selection.")
            valid_count = Supplier.objects.filter(id__in=preferred_supplier_ids).count()
            if valid_count != len(preferred_supplier_ids):
                return HttpResponseBadRequest("One or more selected suppliers are invalid.")

        if supplier_strategy == "unregistered":
            if not unregistered_supplier_name:
                return HttpResponseBadRequest("Enter an unregistered supplier name.")
            _get_or_create_unregistered_supplier(unregistered_supplier_name)
            preferred_supplier_ids = []

        product_ids = request.POST.getlist("product_id[]")
        branch_ids = request.POST.getlist("branch_id[]")
        packet_values = request.POST.getlist("packets[]")

        # Backward compatibility for single-line submissions.
        if not product_ids and request.POST.get("product_id"):
            product_ids = [request.POST.get("product_id")]
            packet_values = [request.POST.get("packets")]
            if can_select_branch:
                branch_ids = [request.POST.get("branch_id")]

        if not product_ids or not packet_values:
            return HttpResponseBadRequest("Add at least one valid order line.")

        if len(product_ids) != len(packet_values):
            return HttpResponseBadRequest("Incomplete order lines were submitted.")

        if can_select_branch and len(branch_ids) != len(product_ids):
            return HttpResponseBadRequest("Each order line must include a target branch.")

        if not can_select_branch and not active_branch:
            return HttpResponseBadRequest("You are not assigned to a valid branch.")

        parsed_lines = []
        for index, (product_id_raw, packets_raw) in enumerate(zip(product_ids, packet_values), start=1):
            product_id_raw = str(product_id_raw or "").strip()
            packets_raw = str(packets_raw or "").strip()

            if not product_id_raw and not packets_raw:
                continue
            if not product_id_raw:
                return HttpResponseBadRequest(f"Select a product on line {index}.")

            if not packets_raw:
                return HttpResponseBadRequest(f"Enter packets to order on line {index}.")

            try:
                packets = int(packets_raw)
            except (TypeError, ValueError):
                return HttpResponseBadRequest(f"Packets must be a whole number on line {index}.")

            if packets <= 0:
                return HttpResponseBadRequest(f"Packets must be greater than zero on line {index}.")

            if can_select_branch:
                branch_id_raw = str(branch_ids[index - 1] or "").strip()
                if not branch_id_raw:
                    return HttpResponseBadRequest(f"Select a branch on line {index}.")
            else:
                branch_id_raw = str(active_branch.id)

            try:
                product_id = int(product_id_raw)
                branch_id = int(branch_id_raw)
            except (TypeError, ValueError):
                return HttpResponseBadRequest(f"Invalid product/branch values on line {index}.")

            parsed_lines.append((product_id, branch_id, packets))

        if not parsed_lines:
            return HttpResponseBadRequest("Add at least one valid order line.")

        product_map = Product.objects.filter(pk__in={line[0] for line in parsed_lines}, is_active=True).in_bulk()
        if len(product_map) != len({line[0] for line in parsed_lines}):
            return HttpResponseBadRequest("One or more selected products are invalid or inactive.")

        if can_select_branch:
            branch_map = Branch.objects.filter(pk__in={line[1] for line in parsed_lines}, is_active=True).in_bulk()
            if len(branch_map) != len({line[1] for line in parsed_lines}):
                return HttpResponseBadRequest("One or more selected branches are invalid or inactive.")
        else:
            branch_map = {active_branch.id: active_branch}

        pending_packets_map = _pending_packets_by_product_branch({line[0] for line in parsed_lines})
        grouped_lines = defaultdict(int)
        for product_id, branch_id, packets in parsed_lines:
            grouped_lines[(product_id, branch_id)] += packets

        for (product_id, branch_id), requested_packets in grouped_lines.items():
            product = product_map[product_id]
            branch = branch_map[branch_id]
            pack_quantity = max(int(product.pack_quantity or 1), 1)
            max_stock_units = max(int(product.max_stock or 1), 1)
            current_branch_units = max(int(product.current_stock(branch) or 0), 0)
            pending_packets_for_branch = pending_packets_map.get((product_id, branch.name), 0)
            pending_branch_units = pending_packets_for_branch * pack_quantity
            available_units = max(max_stock_units - current_branch_units - pending_branch_units, 0)
            max_additional_packets = available_units // pack_quantity

            if requested_packets > max_additional_packets:
                return HttpResponseBadRequest(
                    (
                        f"{product.name} for {branch.name} can only accept {max_additional_packets} more packet(s) "
                        f"without exceeding max stock ({max_stock_units} units). "
                        f"Current stock: {current_branch_units} unit(s), open ordered: {pending_branch_units} unit(s)."
                    )
                )

        reorder_ids = set()
        unregistered_supplier = _get_or_create_unregistered_supplier(unregistered_supplier_name) if supplier_strategy == "unregistered" else None

        with transaction.atomic():
            for (product_id, branch_id), packets in grouped_lines.items():
                product = product_map[product_id]
                branch = branch_map[branch_id]

                reorder = AutoReorderRequest.objects.create(
                    product=product,
                    target_stock_level=max(int(product.max_stock or 1), 1),
                    current_stock_snapshot=product.current_stock(),
                    requested_quantity=packets,
                    remaining_quantity=packets,
                    origin=AutoReorderRequest.ORIGIN_MANUAL,
                    preferred_supplier_ids=preferred_supplier_ids,
                    unregistered_supplier_name=unregistered_supplier_name if supplier_strategy == "unregistered" else "",
                    branch_requirements={branch.name: packets},
                    status=AutoReorderRequest.STATUS_OPEN,
                )
                if supplier_strategy == "unregistered" and unregistered_supplier:
                    now = timezone.now()
                    SupplierReorderRequest.objects.create(
                        reorder_request=reorder,
                        supplier=unregistered_supplier,
                        priority=unregistered_supplier.priority or 9999,
                        requested_quantity=packets,
                        fulfilled_quantity=packets,
                        status=SupplierReorderRequest.STATUS_ACCEPTED,
                        expires_at=now + timedelta(seconds=3600),
                        responded_at=now,
                        emailed_at=now,
                    )
                    reorder.remaining_quantity = 0
                    reorder.status = AutoReorderRequest.STATUS_FULFILLED
                    reorder.completed_at = now
                    reorder.save(update_fields=["remaining_quantity", "status", "completed_at", "updated_at"])
                else:
                    reorder_ids.add(reorder.id)

        for reorder_id in sorted(reorder_ids):
            notify_next_supplier.delay(reorder_id)

        line_count = len(grouped_lines)
        if request.htmx:
            return _with_hx_trigger(
                HttpResponse(""),
                "stock-action-success",
                {
                    "message": (
                        f"Manual order created. "
                        "Supplier notifications sent."
                    )
                },
            )
        messages.success(
            request,
            f"Manual order created. Supplier notifications were queued.",
        )
        return HttpResponse("")

    products = list(Product.objects.filter(is_active=True).order_by("name"))
    product_ids = [product.id for product in products]
    pending_packets_map = _pending_packets_by_product_branch(product_ids)
    branch_scope = list(all_branches) if can_select_branch else ([active_branch] if active_branch else [])
    capacity_by_product_branch = {}

    for product in products:
        product.max_packets_allowed = _max_packets_allowed_for_product(product)
        pack_quantity = max(int(product.pack_quantity or 1), 1)
        max_stock_units = max(int(product.max_stock or 1), 1)
        branch_capacity = {}

        for branch in branch_scope:
            current_branch_units = max(int(product.current_stock(branch) or 0), 0)
            pending_packets_for_branch = pending_packets_map.get((product.id, branch.name), 0)
            pending_branch_units = pending_packets_for_branch * pack_quantity
            available_units = max(max_stock_units - current_branch_units - pending_branch_units, 0)
            max_additional_packets = available_units // pack_quantity
            branch_capacity[str(branch.id)] = {
                "max_packets": int(max_additional_packets),
                "available_units": int(available_units),
                "current_units": int(current_branch_units),
                "pending_units": int(pending_branch_units),
            }

        capacity_by_product_branch[str(product.id)] = branch_capacity

    try:
        normalized_active_branch_id = int(active_branch_id) if active_branch_id else None
    except (TypeError, ValueError):
        normalized_active_branch_id = None
    return render(request, "products/partials/_create_order_form.html", {
        "products": products,
        "suppliers": Supplier.objects.order_by("priority", "name"),
        "all_branches": all_branches,
        "active_branch": active_branch,
        "active_branch_id": normalized_active_branch_id,
        "can_select_branch": can_select_branch,
        "manual_order_capacity_json": json.dumps(capacity_by_product_branch),
    })


def _can_manage_catalog(user):
    return user.is_superuser or user.can_manage_users()


def _can_manage_suppliers(user):
    """Only super admin role can create, edit, or delete suppliers."""
    return user.is_superuser or getattr(user, "role", "") == "super_admin"


def _can_select_products_branch(user):
    return user.is_superuser or getattr(user, "role", "") == "super_admin"


def _can_edit_receive_packets(user):
    return True


def _resolve_products_branch(request):
    selected_branch_id = (
        request.GET.get("branch")
        or request.GET.get("branch_id")
        or request.POST.get("branch_id")
        or request.POST.get("branch")
        or ""
    ).strip()
    can_select_branch = _can_select_products_branch(request.user)
    active_branches = Branch.objects.filter(is_active=True).order_by("name")

    if can_select_branch and selected_branch_id:
        chosen_branch = active_branches.filter(pk=selected_branch_id).first()
        if chosen_branch:
            return chosen_branch

    if request.user.branch_id:
        if can_select_branch:
            active_user_branch = active_branches.filter(pk=request.user.branch_id).first()
            if active_user_branch:
                return active_user_branch
        else:
            user_branch = Branch.objects.filter(pk=request.user.branch_id).first()
            if user_branch:
                return user_branch

    if can_select_branch:
        return active_branches.first()

    return None


def _products_page_context(request, search="", branch=None):
    active_branch = branch if branch is not None else _resolve_products_branch(request)
    can_select_branch = _can_select_products_branch(request.user)
    report_meta = _bulk_upload_report_meta(request)

    return {
        "product_data": _build_product_data(active_branch, search=search),
        "suppliers": Supplier.objects.all(),
        "search": search,
        "active_branch": active_branch,
        "active_branch_id": active_branch.pk if active_branch else "",
        "can_select_branch": can_select_branch,
        "branch_options": Branch.objects.filter(is_active=True).order_by("name") if can_select_branch else [],
        "bulk_upload_report_available": report_meta["available"],
        "bulk_upload_report_failed_count": report_meta["failed_count"],
        "bulk_upload_report_source_filename": report_meta["source_filename"],
        "bulk_upload_report_generated_at": report_meta["generated_at"],
    }


def _redirect_with_branch(route_name, branch):
    url = reverse(route_name)
    if branch:
        url = f"{url}?branch={branch.pk}"
    return redirect(url)


def _build_product_data(branch, search=""):
    products = Product.objects.filter(is_active=True)
    if search:
        products = products.filter(Q(name__icontains=search) | Q(barcode__icontains=search))

    payload = []
    for product in products.order_by("name"):
        stock_qty = product.current_stock(branch) if branch else product.current_stock()
        payload.append({"product": product, "stock": stock_qty})
    return payload


def _render_product_table(request, search="", branch=None):
    ctx = _products_page_context(request, search=search, branch=branch)
    return render(request, "products/partials/_product_table.html", ctx)


def _rows_oob_response(request, products, branch=None, include_messages=True):
    unique_products = []
    seen_ids = set()
    for product in products:
        if not product or product.pk in seen_ids:
            continue
        seen_ids.add(product.pk)
        unique_products.append(product)
    payload = render_to_string("partials/_messages.html", {"oob": True}, request=request) if include_messages else ""

    for product in unique_products:
        stock_qty = product.current_stock(branch) if branch else product.current_stock()
        payload += render_to_string(
            "products/partials/_product_row.html",
            {"p": product, "stock": stock_qty, "oob": True, "active_branch_id": branch.pk if branch else ""},
            request=request,
        )

    return HttpResponse(payload)


def _with_hx_trigger(response, event_name, detail):
    if not response:
        return response

    triggers = {}
    existing = response.headers.get("HX-Trigger")
    if existing:
        try:
            parsed = json.loads(existing)
            if isinstance(parsed, dict):
                triggers.update(parsed)
        except json.JSONDecodeError:
            pass

    triggers[event_name] = detail
    response["HX-Trigger"] = json.dumps(triggers)
    return response


@login_required
def product_list_view(request):
    search = request.GET.get("q", "").strip()
    active_branch = _resolve_products_branch(request)
    ctx = _products_page_context(request, search=search, branch=active_branch)

    if request.htmx:
        return render(request, "products/partials/_product_table.html", ctx)

    return render(request, "products/product_list.html", ctx)


@login_required
def stock_list_view(request):
    search = request.GET.get("q", "").strip()
    active_branch = _resolve_products_branch(request)
    ctx = _products_page_context(request, search=search, branch=active_branch)

    if request.htmx:
        return render(request, "products/partials/_product_table.html", ctx)
    return render(request, "products/stock_list.html", ctx)


@login_required
def product_create_view(request):
    if not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return redirect("product-list")

    active_branch = _resolve_products_branch(request)

    if request.method == "POST":
        form = ProductForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Product created successfully.")
            return _render_product_table(request, branch=active_branch)
        return render(
            request,
            "products/partials/_product_form.html",
            {"form": form, "active_branch_id": active_branch.pk if active_branch else ""},
        )

    form = ProductForm()
    return render(
        request,
        "products/partials/_product_form.html",
        {"form": form, "active_branch_id": active_branch.pk if active_branch else ""},
    )


@login_required
def product_edit_view(request, pk):
    if not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return redirect("product-list")

    active_branch = _resolve_products_branch(request)
    product = get_object_or_404(Product, pk=pk)
    if request.method == "POST":
        form = ProductForm(request.POST, instance=product)
        if form.is_valid():
            form.save()
            messages.success(request, "Product updated.")
            return _render_product_table(request, branch=active_branch)
        return render(
            request,
            "products/partials/_product_form.html",
            {"form": form, "product": product, "active_branch_id": active_branch.pk if active_branch else ""},
        )
    form = ProductForm(instance=product)
    return render(
        request,
        "products/partials/_product_form.html",
        {"form": form, "product": product, "active_branch_id": active_branch.pk if active_branch else ""},
    )


@login_required
def product_detail_view(request, pk):
    product = get_object_or_404(Product, pk=pk)
    active_branch = _resolve_products_branch(request)
    stock_qty = product.current_stock(active_branch) if active_branch else product.current_stock()

    return render(
        request,
        "products/partials/_product_detail.html",
        {
            "product": product,
            "stock": stock_qty,
            "active_branch": active_branch,
        },
    )


@login_required
def receive_stock_view(request):
    if request.user.get_role_display() == "Cashier":
        messages.error(request, "Permission denied.")
        return redirect("/")

    can_select_branch = _can_select_products_branch(request.user)
    can_edit_quantity = _can_edit_receive_packets(request.user)
    active_branch = _resolve_products_branch(request)

    if request.method == "POST":
        if request.POST.get("branch_id") and can_select_branch:
            active_branch = get_object_or_404(Branch, pk=request.POST["branch_id"])

        changed_products = []
        requeue_reorder_ids = set()
        purchase_id_for_confirmation = None

        if not active_branch:
            messages.error(request, "You are not assigned to a branch.")
            return _rows_oob_response(request, changed_products, branch=active_branch) if request.htmx else _redirect_with_branch("stock-list", active_branch)

        try:
            with transaction.atomic():
                supplier = get_object_or_404(Supplier, pk=request.POST.get("supplier_id"))
                invoice_number = request.POST.get("invoice_number", "").strip()
                request_ids = request.POST.getlist("reorder_request_id[]")
                quantities = request.POST.getlist("quantity[]")
                cost_prices = request.POST.getlist("cost_price[]")
                selling_prices = request.POST.getlist("selling_price[]")

                if not invoice_number:
                    raise ValueError("Invoice number is required.")
                if not request_ids:
                    raise ValueError("No pending items found or selected to receive.")
                if len(request_ids) != len(quantities):
                    raise ValueError("Incomplete stock lines were submitted.")
                if len(request_ids) != len(cost_prices) or len(request_ids) != len(selling_prices):
                    raise ValueError("Incomplete pricing lines were submitted.")

                existing_invoice = (
                    Purchase.objects.select_for_update()
                    .filter(
                        supplier=supplier,
                        branch=active_branch,
                        invoice_number__iexact=invoice_number,
                    )
                    .order_by("-created_at")
                    .first()
                )
                if existing_invoice:
                    raise ValueError(
                        (
                            f"Invoice {invoice_number} for {supplier.name} has already been received for "
                            f"{active_branch.name}. Use a different invoice number or review the existing entry."
                        )
                    )

                total_amount = Decimal("0.00")
                cleaned_lines = []
                zero_quantity_lines = []
                for req_id, qty, cost_raw, selling_raw in zip(request_ids, quantities, cost_prices, selling_prices):
                    try:
                        qty_int = int(qty)
                    except (TypeError, ValueError):
                        raise ValueError("Quantity must be a valid whole number.")
                    if qty_int < 0:
                        raise ValueError("Quantity cannot be negative.")

                    sup_req = (
                        SupplierReorderRequest.objects.select_for_update()
                        .select_related("reorder_request__product")
                        .get(pk=req_id)
                    )
                    if sup_req.supplier_id != supplier.id:
                        raise ValueError("One or more selected reorder lines do not belong to the selected supplier.")
                    product = sup_req.reorder_request.product

                    max_allowed_packets = sup_req.pending_quantity
                    if max_allowed_packets <= 0:
                        raise ValueError(f"{product.name} has no pending quantity to receive.")

                    branch_pending_packets = _branch_pending_packets(sup_req, active_branch)
                    max_allowed_packets_for_branch = min(sup_req.pending_quantity, branch_pending_packets)
                    if max_allowed_packets_for_branch <= 0:
                        raise ValueError(
                            f"{sup_req.reorder_request.product.name} is not pending for {active_branch.name}."
                        )

                    if qty_int == 0:
                        zero_quantity_lines.append((sup_req, max_allowed_packets))
                        continue

                    try:
                        cost_dec = Decimal(str(cost_raw).strip())
                    except (TypeError, ValueError, InvalidOperation):
                        raise ValueError(f"Invalid cost price for {product.name}.")

                    try:
                        selling_dec = Decimal(str(selling_raw).strip())
                    except (TypeError, ValueError, InvalidOperation):
                        raise ValueError(f"Invalid selling price for {product.name}.")

                    if cost_dec <= 0:
                        raise ValueError(f"Cost price must be greater than 0 for {product.name}.")

                    pack_quantity = max(int(product.pack_quantity or 1), 1)
                    cost_per_unit = cost_dec / Decimal(pack_quantity)
                    min_selling_price = cost_per_unit * Decimal("1.33")
                    if selling_dec < min_selling_price:
                        raise ValueError(
                            f"Selling price for {product.name} must be at least {min_selling_price.quantize(Decimal('0.01'))} (33% above unit cost)."
                        )

                    if not can_edit_quantity and qty_int != max_allowed_packets_for_branch:
                        raise ValueError(
                            f"QTY (PACKETS) for {sup_req.reorder_request.product.name} can only be edited by super admin."
                        )
                    if qty_int > max_allowed_packets_for_branch:
                        raise ValueError(
                            f"Cannot receive more than {max_allowed_packets_for_branch} packet(s) for {sup_req.reorder_request.product.name} in {active_branch.name}."
                        )

                    actual_units_received = qty_int * product.pack_quantity
                    total_amount += qty_int * cost_dec
                    cleaned_lines.append((sup_req, product, qty_int, actual_units_received, cost_dec, cost_per_unit, selling_dec))

                for sup_req, reverted_packets in zero_quantity_lines:
                    if reverted_packets <= 0:
                        continue

                    original_fulfilled = _to_non_negative_int(sup_req.fulfilled_quantity)
                    original_received = _to_non_negative_int(sup_req.received_quantity)
                    updated_fulfilled = max(original_fulfilled - reverted_packets, original_received)
                    if updated_fulfilled == original_fulfilled:
                        continue

                    sup_req.fulfilled_quantity = updated_fulfilled
                    if sup_req.fulfilled_quantity <= 0:
                        sup_req.status = SupplierReorderRequest.STATUS_REJECTED
                    elif sup_req.fulfilled_quantity < sup_req.requested_quantity:
                        sup_req.status = SupplierReorderRequest.STATUS_PARTIAL
                    else:
                        sup_req.status = SupplierReorderRequest.STATUS_ACCEPTED
                    sup_req.save(update_fields=["fulfilled_quantity", "status", "updated_at"])

                    reorder_request = (
                        AutoReorderRequest.objects.select_for_update()
                        .filter(pk=sup_req.reorder_request_id)
                        .first()
                    )
                    if not reorder_request:
                        continue

                    requested_packets = _to_non_negative_int(reorder_request.requested_quantity)
                    restored_packets = max(original_fulfilled - updated_fulfilled, 0)
                    new_remaining = min(requested_packets, _to_non_negative_int(reorder_request.remaining_quantity) + restored_packets)
                    update_fields = []
                    if reorder_request.remaining_quantity != new_remaining:
                        reorder_request.remaining_quantity = new_remaining
                        update_fields.append("remaining_quantity")
                    if new_remaining > 0 and reorder_request.status != AutoReorderRequest.STATUS_OPEN:
                        reorder_request.status = AutoReorderRequest.STATUS_OPEN
                        update_fields.append("status")
                    if new_remaining > 0 and reorder_request.completed_at is not None:
                        reorder_request.completed_at = None
                        update_fields.append("completed_at")
                    if update_fields:
                        reorder_request.save(update_fields=[*update_fields, "updated_at"])
                    if reorder_request.remaining_quantity > 0:
                        requeue_reorder_ids.add(reorder_request.id)

                if not cleaned_lines:
                    for reorder_id in sorted(requeue_reorder_ids):
                        notify_next_supplier.delay(reorder_id)
                    if request.htmx:
                        response = _rows_oob_response(
                            request,
                            changed_products,
                            branch=active_branch,
                            include_messages=False,
                        )
                        return _with_hx_trigger(
                            response,
                            "stock-action-success",
                            {"message": "Receive completed with zero quantity. No stock was added."},
                        )
                    messages.success(request, "Receive completed with zero quantity. No stock was added.")
                    return _redirect_with_branch("stock-list", active_branch)

                purchase = Purchase.objects.create(
                    supplier=supplier,
                    branch=active_branch,
                    invoice_number=invoice_number,
                    total_amount=total_amount,
                    created_by=request.user,
                )
                purchase_id_for_confirmation = purchase.id

                for sup_req, product, qty_int, actual_units_received, cost_per_packet_dec, cost_per_unit_dec, selling_dec in cleaned_lines:
                    changed_products.append(product)

                    sup_req.received_quantity += qty_int
                    sup_req.save(update_fields=["received_quantity", "updated_at"])

                    reorder_request = sup_req.reorder_request
                    branch_requirements = reorder_request.branch_requirements or {}
                    if (
                        active_branch
                        and isinstance(branch_requirements, dict)
                        and active_branch.name in branch_requirements
                    ):
                        remaining_branch_packets = _to_non_negative_int(branch_requirements.get(active_branch.name)) - qty_int
                        branch_requirements[active_branch.name] = max(remaining_branch_packets, 0)
                        reorder_request.branch_requirements = branch_requirements
                        reorder_request.save(update_fields=["branch_requirements", "updated_at"])

                    if product.cost_price != cost_per_packet_dec or product.unit_price != selling_dec:
                        product.cost_price = cost_per_packet_dec
                        product.unit_price = selling_dec
                        product.save(update_fields=["cost_price", "unit_price", "updated_at"])

                    PurchaseItem.objects.create(
                        purchase=purchase,
                        product=product,
                        quantity=actual_units_received,
                        unit_cost=cost_per_unit_dec,
                    )

                    stock, _ = Stock.objects.get_or_create(
                        product=product,
                        branch=active_branch,
                        defaults={"quantity": 0},
                    )
                    stock.quantity += actual_units_received
                    stock.save(update_fields=["quantity", "updated_at"])

                    StockMovement.objects.create(
                        product=product,
                        branch=active_branch,
                        movement_type="purchase",
                        quantity=actual_units_received,
                        reference=invoice_number,
                        created_by=request.user,
                    )

            if purchase_id_for_confirmation and getattr(settings, "SUPPLIER_RECEIPT_EMAIL_ENABLED", False):
                try:
                    send_purchase_confirmation_to_supplier.delay(purchase_id_for_confirmation)
                except Exception:
                    logger.exception(
                        "Failed to queue supplier purchase confirmation task",
                        extra={"purchase_id": purchase_id_for_confirmation},
                    )

            for reorder_id in sorted(requeue_reorder_ids):
                notify_next_supplier.delay(reorder_id)

            if request.htmx:
                response = _rows_oob_response(
                    request,
                    changed_products,
                    branch=active_branch,
                    include_messages=False,
                )
                return _with_hx_trigger(
                    response,
                    "stock-action-success",
                    {"message": "Stock received successfully."},
                )
                # enable below message for supplier confirmation
            # messages.success(request, "Stock received successfully. Supplier Receipt confirmation has been sent.")
            # return _redirect_with_branch("stock-list", active_branch)

        except Exception as exc:
            messages.error(request, f"Error: {exc}")
            return _rows_oob_response(request, changed_products, branch=active_branch) if request.htmx else _redirect_with_branch("stock-list", active_branch)

    suppliers = Supplier.objects.all()
    products = Product.objects.filter(is_active=True)
    all_branches = Branch.objects.filter(is_active=True).order_by("name") if can_select_branch else None
    template_name = "products/partials/_receive_form.html" if request.htmx else "products/receive_stock.html"
    return render(
        request,
        template_name,
        {
            "suppliers": suppliers,
            "products": products,
            "all_branches": all_branches,
            "active_branch": active_branch,
            "active_branch_id": active_branch.pk if active_branch else "",
        },
    )


@login_required
def adjust_stock_view(request, pk):
    product = get_object_or_404(Product, pk=pk)
    branch = _resolve_products_branch(request)

    if not request.user.can_adjust_stock() and not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return _rows_oob_response(request, [product], branch=branch) if request.htmx else _redirect_with_branch("stock-list", branch)

    if not branch:
        messages.error(request, "Select or assign a branch first.")
        return _rows_oob_response(request, [product], branch=branch) if request.htmx else _redirect_with_branch("stock-list", branch)

    if request.method == "POST":
        form = AdjustStockForm(request.POST)
        if form.is_valid():
            try:
                with transaction.atomic():
                    new_qty = form.cleaned_data["new_quantity"]

                    stock, _ = Stock.objects.get_or_create(product=product, branch=branch, defaults={"quantity": 0})
                    adjustment = new_qty - stock.quantity
                    stock.quantity = new_qty
                    stock.save(update_fields=["quantity", "updated_at"])

                    StockMovement.objects.create(
                        product=product,
                        branch=branch,
                        movement_type="adjustment",
                        quantity=adjustment,
                        created_by=request.user,
                    )

                messages.success(request, f"Stock adjusted for {product.name}.")
                return _rows_oob_response(request, [product], branch=branch) if request.htmx else _redirect_with_branch("stock-list", branch)
            except Exception as exc:
                messages.error(request, f"Error: {exc}")
                return _rows_oob_response(request, [product], branch=branch) if request.htmx else _redirect_with_branch("stock-list", branch)

        template_name = "products/partials/_adjust_form.html" if request.htmx else "products/adjust_stock.html"
        return render(
            request,
            template_name,
            {
                "form": form,
                "product": product,
                "current_stock": product.current_stock(branch),
                "active_branch": branch,
                "active_branch_id": branch.pk if branch else "",
            },
        )

    form = AdjustStockForm()
    template_name = "products/partials/_adjust_form.html" if request.htmx else "products/adjust_stock.html"
    return render(
        request,
        template_name,
        {
            "form": form,
            "product": product,
            "current_stock": product.current_stock(branch),
            "active_branch": branch,
            "active_branch_id": branch.pk if branch else "",
        },
    )


@login_required
def transfer_stock_view(request):
    if request.user.get_role_display() == "Cashier":
        messages.error(request, "Permission denied.")
        return redirect("/")

    from_branch = _resolve_products_branch(request)
    template_name = "products/partials/_transfer_form.html" if request.htmx else "products/transfer_stock.html"
    
    if not from_branch:
        messages.error(request, "You are not assigned to any branch.")
        if request.method == "GET":
            return render(request, template_name, {"blocked": True})
        return _rows_oob_response(request, [], branch=from_branch) if request.htmx else _redirect_with_branch("stock-list", from_branch)

    destination_branches = Branch.objects.filter(is_active=True).exclude(pk=from_branch.pk).order_by("name")

    if request.method == "POST":
        if request.POST.get("from_branch_id"):
            from_branch = get_object_or_404(Branch, pk=request.POST["from_branch_id"])

        destination_branches = Branch.objects.filter(is_active=True).exclude(pk=from_branch.pk).order_by("name")
        form = TransferStockForm(request.POST)
        form.fields["to_branch"].queryset = Branch.objects.filter(is_active=True).order_by("name")
        changed_products = []

        if form.is_valid():
            to_branch = form.cleaned_data["to_branch"]
            if from_branch and to_branch and from_branch.pk == to_branch.pk:
                form.add_error("to_branch", "Destination branch must be different from source branch.")
                all_branches = Branch.objects.filter(is_active=True).order_by("name")
                products = Product.objects.filter(is_active=True).order_by("name")
                return render(
                    request,
                    template_name,
                    {
                        "form": form,
                        "products": products,
                        "all_branches": all_branches,
                        "branches": destination_branches,
                        "destination_branches": destination_branches,
                        "selected_to_branch_id": request.POST.get("to_branch", ""),
                        "active_branch": from_branch,
                        "active_branch_id": from_branch.pk if from_branch else "",
                    },
                )
            notes = form.cleaned_data.get("notes", "")
            product_ids = request.POST.getlist("product_id[]")
            quantities = request.POST.getlist("quantity[]")
            if not product_ids:
                messages.error(request, "Add at least one product line.")
                return _rows_oob_response(request, changed_products, branch=from_branch) if request.htmx else _redirect_with_branch("stock-list", from_branch)
            if len(product_ids) != len(quantities):
                messages.error(request, "Incomplete transfer lines were submitted.")
                return _rows_oob_response(request, changed_products, branch=from_branch) if request.htmx else _redirect_with_branch("stock-list", from_branch)

            try:
                with transaction.atomic():
                    transfer = Transfer.objects.create(
                        from_branch=from_branch,
                        to_branch=to_branch,
                        created_by=request.user,
                        notes=notes,
                    )

                    for pid, qty in zip(product_ids, quantities):
                        qty_int = int(qty)
                        if qty_int <= 0:
                            raise ValueError("Transfer quantity must be greater than zero.")

                        product = get_object_or_404(Product, pk=pid)
                        from_stock, _ = Stock.objects.get_or_create(
                            product=product,
                            branch=from_branch,
                            defaults={"quantity": 0},
                        )
                        if from_stock.quantity < qty_int:
                            raise ValueError(f"Insufficient stock for {product.name}.")

                        to_stock, _ = Stock.objects.get_or_create(
                            product=product,
                            branch=to_branch,
                            defaults={"quantity": 0},
                        )

                        TransferItem.objects.create(
                            transfer=transfer,
                            product=product,
                            quantity=qty_int,
                        )

                        from_stock.quantity -= qty_int
                        from_stock.save(update_fields=["quantity", "updated_at"])
                        to_stock.quantity += qty_int
                        to_stock.save(update_fields=["quantity", "updated_at"])

                        StockMovement.objects.create(
                            product=product,
                            branch=from_branch,
                            movement_type="transfer_out",
                            quantity=-qty_int,
                            reference=f"TRF-{transfer.id}",
                            notes=notes,
                            created_by=request.user,
                        )
                        StockMovement.objects.create(
                            product=product,
                            branch=to_branch,
                            movement_type="transfer_in",
                            quantity=qty_int,
                            reference=f"TRF-{transfer.id}",
                            notes=notes,
                            created_by=request.user,
                        )

                        changed_products.append(product)

                if request.htmx:
                    response = _rows_oob_response(
                        request,
                        changed_products,
                        branch=from_branch,
                        include_messages=False,
                    )
                    return _with_hx_trigger(
                        response,
                        "stock-action-success",
                        {"message": f"Transfer to {to_branch.name} completed."},
                    )
                messages.success(request, f"Transfer to {to_branch.name} completed.")
                return _redirect_with_branch("stock-list", from_branch)
            except Exception as exc:
                messages.error(request, f"Error: {exc}")
                return _rows_oob_response(request, changed_products, branch=from_branch) if request.htmx else _redirect_with_branch("stock-list", from_branch)

        all_branches = Branch.objects.filter(is_active=True).order_by("name")
        products = Product.objects.filter(is_active=True).order_by("name")
        return render(
            request,
            template_name,
            {
                "form": form,
                "products": products,
                "all_branches": all_branches,
                "branches": destination_branches,
                "destination_branches": destination_branches,
                "selected_to_branch_id": request.POST.get("to_branch", ""),
                "active_branch": from_branch,
                "active_branch_id": from_branch.pk if from_branch else "",
            },
        )

    form = TransferStockForm()
    form.fields["to_branch"].queryset = destination_branches
    all_branches = Branch.objects.filter(is_active=True).order_by("name")
    products = Product.objects.filter(is_active=True).order_by("name")
    return render(
        request,
        template_name,
        {
            "form": form,
            "products": products,
            "all_branches": all_branches,
            "branches": destination_branches,
            "destination_branches": destination_branches,
            "selected_to_branch_id": "",
            "active_branch": from_branch,
            "active_branch_id": from_branch.pk if from_branch else "",
        },
    )


@login_required
@require_GET
def transfer_destination_branches_view(request):
    from_branch_id = (request.GET.get("from_branch_id") or "").strip()
    selected_to_branch_id = (
        request.GET.get("selected_to_branch_id")
        or request.GET.get("to_branch")
        or ""
    ).strip()

    destination_branches = Branch.objects.filter(is_active=True).order_by("name")
    if from_branch_id:
        destination_branches = destination_branches.exclude(pk=from_branch_id)

    return render(
        request,
        "products/partials/_transfer_destination_field.html",
        {
            "destination_branches": destination_branches,
            "selected_to_branch_id": selected_to_branch_id,
            "to_branch_errors": "",
        },
    )


def _normalized_header_name(value):
    return str(value or "").strip().lower().replace(" ", "_")


def _to_cell_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _row_cell_value(row, index):
    if index is None:
        return None
    try:
        return row[index]
    except (IndexError, TypeError):
        return None


def _parse_decimal_cell(value, field_label):
    text_value = _to_cell_text(value)
    if not text_value:
        raise ValueError(f"Missing {field_label}.")
    try:
        return Decimal(text_value.replace(",", ""))
    except (InvalidOperation, ValueError):
        raise ValueError(f"Invalid {field_label}: {text_value}")


def _parse_int_cell(value, field_label):
    parsed_decimal = _parse_decimal_cell(value, field_label)
    try:
        return int(parsed_decimal)
    except (TypeError, ValueError):
        raise ValueError(f"Invalid {field_label}: {_to_cell_text(value)}")


def _bulk_upload_report_meta(request):
    report = request.session.get(BULK_UPLOAD_FAILED_ROWS_SESSION_KEY) or {}
    rows = report.get("rows") or []
    return {
        "available": bool(rows),
        "failed_count": len(rows),
        "source_filename": report.get("source_filename", ""),
        "generated_at": report.get("generated_at", ""),
    }


@login_required
def product_bulk_template_download_view(request):
    try:
        from openpyxl import Workbook
    except ImportError:
        messages.error(request, "openpyxl is not installed. Please install dependencies and retry.")
        return redirect("product-list")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory Template"
    sheet.append(["CODE", "DESCRIPTION", "PACK_QTY", "INV_TRADE_PRICE", "SELLING_PRICE", "MINIMUM", "MAXIMUM"])
    sheet.append(["1000001", "Paracetamol 500mg", 10, 50.00, 80.00, 20, 200])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="inventory_upload_template.xlsx"'
    return response


@login_required
def product_bulk_upload_view(request):
    active_branch = _resolve_products_branch(request)

    if request.method == "POST":
        form = ProductBulkUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(
                request,
                "products/partials/_bulk_upload_form.html",
                {"form": form, "active_branch_id": active_branch.pk if active_branch else ""},
            )

        try:
            from openpyxl import load_workbook
        except ImportError:
            messages.error(request, "openpyxl is not installed. Please install dependencies and retry.")
            return _render_product_table(request, branch=active_branch)

        upload = form.cleaned_data["file"]
        workbook = load_workbook(upload, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            messages.error(request, "The uploaded sheet is empty.")
            return _render_product_table(request, branch=active_branch)

        headers = [_normalized_header_name(col) for col in rows[0]]
        required_headers = {"code", "description", "pack_qty", "inv_trade_price", "selling_price", "minimum", "maximum"}
        if not required_headers.issubset(set(headers)):
            messages.error(
                request,
                "Missing required columns. Required: CODE, DESCRIPTION, PACK_QTY, INV_TRADE_PRICE, "
                "SELLING_PRICE, MINIMUM, MAXIMUM.",
            )
            return _render_product_table(request, branch=active_branch)

        h = {header: idx for idx, header in enumerate(headers)}
        created = 0
        updated = 0
        skipped = 0
        failed_rows = []

        for row_index, row in enumerate(rows[1:], start=2):
            if row is None or all(cell in (None, "") for cell in row):
                continue

            barcode = _to_cell_text(_row_cell_value(row, h.get("code")))
            description = _to_cell_text(_row_cell_value(row, h.get("description")))
            pack_qty_text = _to_cell_text(_row_cell_value(row, h.get("pack_qty")))
            trade_price_text = _to_cell_text(_row_cell_value(row, h.get("inv_trade_price")))
            selling_price_text = _to_cell_text(_row_cell_value(row, h.get("selling_price")))
            minimum_text = _to_cell_text(_row_cell_value(row, h.get("minimum")))
            maximum_text = _to_cell_text(_row_cell_value(row, h.get("maximum")))

            try:
                pack_quantity = _parse_int_cell(pack_qty_text, "PACK_QTY")
                cost_price = _parse_decimal_cell(trade_price_text, "INV_TRADE_PRICE")
                unit_price = _parse_decimal_cell(selling_price_text, "SELLING_PRICE")
                reorder_level = _parse_int_cell(minimum_text, "MINIMUM")
                max_stock = _parse_int_cell(maximum_text, "MAXIMUM")
                if not description:
                    raise ValueError("Missing DESCRIPTION.")
                if pack_quantity < 1 or max_stock < 1 or reorder_level < 0:
                    raise ValueError("PACK_QTY and MAXIMUM must be at least 1, and MINIMUM cannot be negative.")
            except Exception as exc:
                skipped += 1
                reason = str(exc).strip() or "Invalid row data."
                failed_rows.append(
                    {
                        "row_number": row_index,
                        "code": barcode,
                        "description": description,
                        "pack_qty": pack_qty_text,
                        "inv_trade_price": trade_price_text,
                        "selling_price": selling_price_text,
                        "minimum": minimum_text,
                        "maximum": maximum_text,
                        "reason": reason,
                    }
                )
                continue

            name = description
            defaults = {
                "name": name,
                "description": description,
                "unit_price": unit_price,
                "cost_price": cost_price,
                "reorder_level": reorder_level,
                "max_stock": max_stock,
                "pack_quantity": pack_quantity,
                "is_active": True,
            }

            if barcode:
                product, was_created = Product.objects.update_or_create(barcode=barcode, defaults=defaults)
            else:
                product_qs = Product.objects.filter(name__iexact=name)
                if product_qs.exists():
                    product = product_qs.first()
                    for field, value in defaults.items():
                        setattr(product, field, value)
                    product.save()
                    was_created = False
                else:
                    product = Product.objects.create(barcode="", **defaults)
                    was_created = True

            if was_created:
                created += 1
            else:
                updated += 1

        messages.success(
            request,
            f"Bulk upload complete. Created: {created}, Updated: {updated}, Skipped: {skipped}.",
        )
        if failed_rows:
            request.session[BULK_UPLOAD_FAILED_ROWS_SESSION_KEY] = {
                "generated_at": timezone.now().strftime("%Y-%m-%d %H:%M:%S"),
                "source_filename": upload.name,
                "rows": failed_rows,
            }
            request.session.modified = True
            messages.warning(
                request,
                f"{len(failed_rows)} row(s) were skipped. Use 'Download Failed Rows CSV' on the Products page.",
            )
        elif BULK_UPLOAD_FAILED_ROWS_SESSION_KEY in request.session:
            del request.session[BULK_UPLOAD_FAILED_ROWS_SESSION_KEY]
            request.session.modified = True
        return _render_product_table(request, branch=active_branch)

    return render(
        request,
        "products/partials/_bulk_upload_form.html",
        {"form": ProductBulkUploadForm(), "active_branch_id": active_branch.pk if active_branch else ""},
    )


@login_required
@require_GET
def product_bulk_upload_failed_rows_download_view(request):
    report = request.session.get(BULK_UPLOAD_FAILED_ROWS_SESSION_KEY) or {}
    failed_rows = report.get("rows") or []
    if not failed_rows:
        messages.error(request, "No failed rows report found. Upload a bulk file first.")
        return redirect("product-list")

    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="bulk_upload_failed_rows_{timestamp}.csv"'

    writer = csv.writer(response)
    writer.writerow(
        [
            "row_number",
            "code",
            "DESCRIPTION",
            "PACK_QTY",
            "INV_TRADE_PRICE",
            "SELLING_PRICE",
            "MINIMUM",
            "MAXIMUM",
            "reason",
        ]
    )
    for item in failed_rows:
        writer.writerow(
            [
                item.get("row_number", ""),
                item.get("code", ""),
                item.get("description", ""),
                item.get("pack_qty", ""),
                item.get("inv_trade_price", ""),
                item.get("selling_price", ""),
                item.get("minimum", ""),
                item.get("maximum", ""),
                item.get("reason", ""),
            ]
        )

    return response


@login_required
def supplier_list_view(request):
    if not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return redirect("dashboard")

    can_edit_suppliers = _can_manage_suppliers(request.user)
    suppliers = Supplier.objects.order_by("name")
    ctx = {
        "suppliers": suppliers,
        "form": SupplierForm(),
        "can_edit_suppliers": can_edit_suppliers,
    }

    if request.htmx:
        return render(request, "products/partials/_supplier_table.html", ctx)
    return render(request, "products/supplier_list.html", ctx)


@login_required
def supplier_create_view(request):
    if not _can_manage_suppliers(request.user):
        messages.error(request, "Permission denied. Only Super Admins can add suppliers.")
        return redirect("supplier-list")

    if request.method == "POST":
        form = SupplierForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Supplier created successfully.")
            suppliers = Supplier.objects.order_by("name")
            return render(
                request,
                "products/partials/_supplier_table.html",
                {"suppliers": suppliers, "can_edit_suppliers": True},
            )
        return render(request, "products/partials/_supplier_form.html", {"form": form})

    return render(request, "products/partials/_supplier_form.html", {"form": SupplierForm()})


@login_required
def supplier_edit_view(request, pk):
    if not _can_manage_suppliers(request.user):
        messages.error(request, "Permission denied. Only Super Admins can edit suppliers.")
        return redirect("supplier-list")

    supplier = get_object_or_404(Supplier, pk=pk)
    if request.method == "POST":
        form = SupplierForm(request.POST, instance=supplier)
        if form.is_valid():
            form.save()
            messages.success(request, "Supplier updated.")
            suppliers = Supplier.objects.order_by("name")
            return render(
                request,
                "products/partials/_supplier_table.html",
                {"suppliers": suppliers, "can_edit_suppliers": True},
            )
        return render(request, "products/partials/_supplier_form.html", {"form": form, "supplier": supplier})

    return render(
        request,
        "products/partials/_supplier_form.html",
        {"form": SupplierForm(instance=supplier), "supplier": supplier},
    )


@login_required
def supplier_delete_view(request, pk):
    if not _can_manage_suppliers(request.user):
        messages.error(request, "Permission denied. Only Super Admins can delete suppliers.")
        return redirect("supplier-list")

    supplier = get_object_or_404(Supplier, pk=pk)

    if request.method == "POST":
        supplier_name = supplier.name
        supplier.delete()
        messages.success(request, f"Supplier '{supplier_name}' deleted successfully.")
        suppliers = Supplier.objects.order_by("name")
        return render(
            request,
            "products/partials/_supplier_table.html",
            {"suppliers": suppliers, "can_edit_suppliers": True},
        )

    # GET — render confirmation modal
    return render(
        request,
        "products/partials/_supplier_delete_confirm.html",
        {"supplier": supplier},
    )


@login_required
def category_list_view(request):
    if not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return redirect("dashboard")

    categories = Category.objects.order_by("name")
    if request.htmx:
        return render(request, "products/partials/_category_table.html", {"categories": categories})
    return render(request, "products/category_list.html", {"categories": categories})


@login_required
def category_create_view(request):
    if not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return redirect("dashboard")

    if request.method == "POST":
        form = CategoryForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Category created successfully.")
            categories = Category.objects.order_by("name")
            return render(request, "products/partials/_category_table.html", {"categories": categories})
        return render(request, "products/partials/_category_form.html", {"form": form})

    return render(request, "products/partials/_category_form.html", {"form": CategoryForm()})


@login_required
def category_edit_view(request, pk):
    if not _can_manage_catalog(request.user):
        messages.error(request, "Permission denied.")
        return redirect("dashboard")

    category = get_object_or_404(Category, pk=pk)
    if request.method == "POST":
        form = CategoryForm(request.POST, instance=category)
        if form.is_valid():
            form.save()
            messages.success(request, "Category updated.")
            categories = Category.objects.order_by("name")
            return render(request, "products/partials/_category_table.html", {"categories": categories})
        return render(request, "products/partials/_category_form.html", {"form": form, "category": category})

    return render(request, "products/partials/_category_form.html", {"form": CategoryForm(instance=category), "category": category})


def supplier_reorder_response_view(request, token):
    primary_request = get_object_or_404(
        SupplierReorderRequest.objects.select_related("supplier", "reorder_request__product"),
        token=token,
    )
    supplier = primary_request.supplier
    now = timezone.now()

    # Expire any pending expired requests for this supplier
    pending_expired = SupplierReorderRequest.objects.filter(
        supplier=supplier,
        status=SupplierReorderRequest.STATUS_PENDING,
        expires_at__lt=now
    ).select_related("reorder_request")

    for locked_request in pending_expired:
        with transaction.atomic():
            locked = SupplierReorderRequest.objects.select_for_update().filter(
                pk=locked_request.pk, status=SupplierReorderRequest.STATUS_PENDING
            ).first()
            if locked:
                locked.status = SupplierReorderRequest.STATUS_EXPIRED
                locked.responded_at = now
                locked.save(update_fields=["status", "responded_at", "updated_at"])
                reorder = locked.reorder_request
                if reorder.status == AutoReorderRequest.STATUS_OPEN and reorder.remaining_quantity > 0:
                    notify_next_supplier.delay(reorder.id)

    pending_requests_qs = SupplierReorderRequest.objects.filter(
        supplier=supplier,
        status=SupplierReorderRequest.STATUS_PENDING,
        expires_at__gte=now
    ).select_related("reorder_request__product").order_by("created_at")
    pending_requests = list(pending_requests_qs)

    branch_product_map = {}
    for pending_request in pending_requests:
        branch_breakdown = _normalized_branch_requirements(
            pending_request.reorder_request.branch_requirements or {}
        )
        pending_request.branch_breakdown = branch_breakdown
        pack_quantity = max(int(pending_request.reorder_request.product.pack_quantity or 1), 1)
        for branch_name, qty in branch_breakdown:
            branch_product_map.setdefault(branch_name, []).append(
                {
                    "product_name": pending_request.reorder_request.product.name,
                    "quantity": qty,
                    "quantity_units": qty * pack_quantity,
                    "req_id": pending_request.id,
                    "requested_quantity": pending_request.requested_quantity,
                }
            )

    branch_groups = _build_branch_groups(branch_product_map)
    primary_request.refresh_from_db()
    primary_request_is_live = (
        primary_request.status == SupplierReorderRequest.STATUS_PENDING and primary_request.expires_at >= now
    )

    if not primary_request_is_live:
        message_type, message = _used_supplier_link_message(primary_request)
        used_branch_groups = _confirmed_branch_groups_for_supplier_request(primary_request)
        pack_quantity = max(int(primary_request.reorder_request.product.pack_quantity or 1), 1)
        reference_request = {
            "product_name": primary_request.reorder_request.product.name,
            "requested_quantity": primary_request.requested_quantity,
            "requested_quantity_units": primary_request.requested_quantity * pack_quantity,
            "fulfilled_quantity": primary_request.fulfilled_quantity,
            "fulfilled_quantity_units": primary_request.fulfilled_quantity * pack_quantity,
            "status_display": primary_request.get_status_display(),
        }
        return render(
            request,
            "products/supplier_reorder_response.html",
            {
                "supplier": supplier,
                "can_respond": False,
                "message_type": message_type,
                "message": message,
                "reference_request": reference_request,
                "branch_groups": used_branch_groups,
                "branch_summary_title": "Confirmed Supply Summary" if used_branch_groups else "Request Summary",
                "branch_summary_subtitle": (
                    "Products you confirmed to supply, grouped by branch"
                    if used_branch_groups
                    else "This request is closed and cannot be edited."
                ),
                "pending_count": 0,
            },
        )

    if request.method == "POST":
        should_escalate_ids = []
        confirmed_branch_product_map = {}

        submitted_branch_payload = defaultdict(dict)
        for row_ref in request.POST.getlist("row_ref"):
            row_ref = (row_ref or "").strip()
            if not row_ref:
                continue

            req_id_raw = (request.POST.get(f"request_id_{row_ref}") or "").strip()
            branch_name = (request.POST.get(f"branch_name_{row_ref}") or "").strip()
            branch_requested_qty = _to_non_negative_int(request.POST.get(f"branch_requested_{row_ref}"))

            try:
                req_id = int(req_id_raw)
            except (TypeError, ValueError):
                continue

            if not branch_name:
                continue

            can_supply = request.POST.get(f"can_supply_{row_ref}") == "yes"
            quantity = None
            quantity_raw = (request.POST.get(f"quantity_{row_ref}") or "").strip()
            if quantity_raw:
                try:
                    quantity = int(quantity_raw)
                except ValueError:
                    quantity = None

            if can_supply:
                if quantity is None or quantity <= 0:
                    quantity = branch_requested_qty
                quantity = min(max(quantity, 0), branch_requested_qty)
            else:
                quantity = 0

            submitted_branch_payload[req_id][branch_name] = quantity

        for req_id, submitted_branches in submitted_branch_payload.items():
            with transaction.atomic():
                locked_req = (
                    SupplierReorderRequest.objects.select_for_update()
                    .select_related("reorder_request__product")
                    .filter(pk=req_id, supplier=supplier, status=SupplierReorderRequest.STATUS_PENDING)
                    .first()
                )
                if not locked_req:
                    continue

                if now > locked_req.expires_at:
                    continue

                reorder = (
                    AutoReorderRequest.objects.select_for_update()
                    .filter(pk=locked_req.reorder_request_id, status=AutoReorderRequest.STATUS_OPEN)
                    .first()
                )
                if not reorder:
                    locked_req.status = SupplierReorderRequest.STATUS_REJECTED
                    locked_req.responded_at = now
                    locked_req.save(update_fields=["status", "responded_at", "updated_at"])
                    continue

                effective_max = min(locked_req.requested_quantity, reorder.remaining_quantity)
                requested_branches = _normalized_branch_requirements(reorder.branch_requirements or {})
                confirmed_branch_allocations = []
                confirmed_quantity = 0

                for branch_name, requested_branch_qty in requested_branches:
                    branch_quantity = min(
                        _to_non_negative_int(submitted_branches.get(branch_name, 0)),
                        requested_branch_qty,
                    )
                    if branch_quantity <= 0:
                        continue
                    confirmed_branch_allocations.append((branch_name, branch_quantity))
                    confirmed_quantity += branch_quantity

                if not requested_branches and effective_max > 0:
                    fallback_quantity = min(_to_non_negative_int(sum(submitted_branches.values())), effective_max)
                    if fallback_quantity > 0:
                        confirmed_branch_allocations = [("Unspecified Branch", fallback_quantity)]
                        confirmed_quantity = fallback_quantity

                if confirmed_quantity > effective_max:
                    remaining_confirmable = effective_max
                    trimmed_allocations = []
                    for branch_name, branch_quantity in confirmed_branch_allocations:
                        if remaining_confirmable <= 0:
                            break
                        adjusted_quantity = min(branch_quantity, remaining_confirmable)
                        if adjusted_quantity <= 0:
                            continue
                        trimmed_allocations.append((branch_name, adjusted_quantity))
                        remaining_confirmable -= adjusted_quantity
                    confirmed_branch_allocations = trimmed_allocations
                    confirmed_quantity = effective_max

                if confirmed_quantity <= 0:
                    locked_req.status = SupplierReorderRequest.STATUS_REJECTED
                    locked_req.fulfilled_quantity = 0
                    locked_req.responded_at = now
                    locked_req.save(update_fields=["status", "fulfilled_quantity", "responded_at", "updated_at"])
                    if reorder.remaining_quantity > 0:
                        should_escalate_ids.append(reorder.id)
                else:
                    locked_req.fulfilled_quantity = confirmed_quantity
                    if confirmed_quantity < locked_req.requested_quantity:
                        locked_req.status = SupplierReorderRequest.STATUS_PARTIAL
                    else:
                        locked_req.status = SupplierReorderRequest.STATUS_ACCEPTED
                    locked_req.responded_at = now
                    locked_req.save(update_fields=["status", "fulfilled_quantity", "responded_at", "updated_at"])

                    reorder.remaining_quantity = max(reorder.remaining_quantity - confirmed_quantity, 0)
                    if reorder.remaining_quantity == 0:
                        reorder.status = AutoReorderRequest.STATUS_FULFILLED
                        reorder.completed_at = now
                    reorder.save(update_fields=["remaining_quantity", "status", "completed_at", "updated_at"])

                    if reorder.remaining_quantity > 0:
                        should_escalate_ids.append(reorder.id)
                    pack_quantity = max(int(locked_req.reorder_request.product.pack_quantity or 1), 1)
                    for branch_name, branch_qty in confirmed_branch_allocations:
                        confirmed_branch_product_map.setdefault(branch_name, []).append(
                            {
                                "product_name": locked_req.reorder_request.product.name,
                                "quantity": branch_qty,
                                "quantity_units": branch_qty * pack_quantity,
                            }
                        )

        for esc_id in set(should_escalate_ids):
            notify_next_supplier.delay(esc_id)

        confirmed_branch_groups = _build_branch_groups(confirmed_branch_product_map)

        return render(
            request,
            "products/supplier_reorder_response.html",
            {
                "supplier": supplier,
                "can_respond": False,
                "message_type": "success",
                "message": "Feedback received successfully. Thank you for your response.",
                "branch_groups": confirmed_branch_groups,
                "branch_summary_title": "Confirmed Supply Summary",
                "branch_summary_subtitle": "Products you confirmed to supply, grouped by branch",
                "pending_count": 0,
            },
        )

    return render(
        request,
        "products/supplier_reorder_response.html",
        {
            "supplier": supplier,
            "pending_requests": pending_requests,
            "can_respond": bool(pending_requests),
            "branch_groups": branch_groups,
            "branch_summary_title": "Branch Summary",
            "branch_summary_subtitle": "All needed products grouped by branch",
            "pending_count": len(pending_requests),
        },
    )


@login_required
def supplier_prioritize_view(request):
    if not getattr(request.user, "can_manage_users", False):
        messages.error(request, "Permission denied.")
        return redirect("/")
        
    if request.method == "POST":
        supplier_ids = request.POST.getlist("supplier_id")
        try:
            with transaction.atomic():
                for idx, s_id in enumerate(supplier_ids):
                    Supplier.objects.filter(id=int(s_id)).update(priority=idx + 1)
            messages.success(request, f"Global supplier priorities updated successfully.")
        except Exception as e:
            messages.error(request, f"Error updating priorities: {e}")
        return redirect("supplier-prioritize")
        
    suppliers = Supplier.objects.all().order_by("priority", "id")
    
    return render(request, "products/supplier_prioritize.html", {
        "suppliers": suppliers,
    })
